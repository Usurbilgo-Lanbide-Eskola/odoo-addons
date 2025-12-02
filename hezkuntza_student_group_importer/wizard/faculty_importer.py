# Copyright 2025 Mikel Arregi Etxaniz - CIFP Usurbil LHII
# License AGPL-3.0 or later (https://www.gnu.org/licenses/agpl.html).

import base64
import io
import logging
from datetime import datetime

from odoo import _, api, fields, models
from odoo.exceptions import UserError, ValidationError

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
except ImportError:
    openpyxl = None

_logger = logging.getLogger(__name__)


class FacultyImporter(models.TransientModel):
    _name = 'faculty.importer'
    _description = 'Faculty Importer Wizard'

    file = fields.Binary(string='Excel File')
    filename = fields.Char(string='Filename')
    override = fields.Boolean(
        string='Override Existing Data',
        default=False,
        help='If checked, existing faculty data will be updated with imported data'
    )
    template_file = fields.Binary(
        string='Template File',
        compute='_compute_template_file',
        readonly=True
    )
    template_filename = fields.Char(
        string='Template Filename',
        default='faculty_import_template.xlsx',
        readonly=True
    )

    @api.depends('template_filename')
    def _compute_template_file(self):
        """Generate Excel template with all speciality codes"""
        for wizard in self:
            if not openpyxl:
                wizard.template_file = False
                continue

            # Create workbook and worksheet
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Faculty Import"

            # Define headers
            headers = [
                'lastname',
                'second_lastname',
                'name',
                'id',
                'email',
                'gender',
                'birthdate',
                'speciality'
            ]

            # Style for headers
            header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
            header_font = Font(bold=True, color='FFFFFF')

            # Write headers
            for col_num, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_num, value=header)
                cell.fill = header_fill
                cell.font = header_font

            # Add example row
            example_data = [
                'Doe',
                'Smith',
                'John',
                '12345678A',
                'john.doe@example.com',
                'o',
                '1980-01-15',
                'inf'
            ]
            for col_num, value in enumerate(example_data, 1):
                ws.cell(row=2, column=col_num, value=value)

            # Get all specialities
            specialities = self.env['hezkuntza.speciality'].search([])
            
            # Create specialities sheet
            ws_spec = wb.create_sheet("Specialities Reference")
            ws_spec.cell(row=1, column=1, value='Code').fill = header_fill
            ws_spec.cell(row=1, column=1).font = header_font
            ws_spec.cell(row=1, column=2, value='Name').fill = header_fill
            ws_spec.cell(row=1, column=2).font = header_font

            for idx, spec in enumerate(specialities, 2):
                ws_spec.cell(row=idx, column=1, value=spec.name)
                ws_spec.cell(row=idx, column=2, value=spec.description)


            # Create gender reference sheet
            ws_gender = wb.create_sheet("Gender Reference")
            ws_gender.cell(row=1, column=1, value='Code').fill = header_fill
            ws_gender.cell(row=1, column=1).font = header_font
            ws_gender.cell(row=1, column=2, value='Description').fill = header_fill
            ws_gender.cell(row=1, column=2).font = header_font

            gender_options = [
                ('m', 'Male'),
                ('f', 'Female'),
                ('o', 'Other/Non-binary'),
            ]
            for idx, (code, desc) in enumerate(gender_options, 2):
                ws_gender.cell(row=idx, column=1, value=code)
                ws_gender.cell(row=idx, column=2, value=desc)

            # Adjust column widths
            for ws_temp in [ws, ws_spec, ws_gender]:
                for column in ws_temp.columns:
                    max_length = 0
                    column_letter = column[0].column_letter
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(cell.value)
                        except:
                            pass
                    adjusted_width = (max_length + 2)
                    ws_temp.column_dimensions[column_letter].width = adjusted_width

            # Save to binary
            output = io.BytesIO()
            wb.save(output)
            wizard.template_file = base64.b64encode(output.getvalue())

    def action_download_template(self):
        """Download the template file"""
        self.ensure_one()
        return {
            'type': 'ir.actions.act_url',
            'url': f'/web/content?model={self._name}&id={self.id}&field=template_file&filename_field=template_filename&download=true',
            'target': 'self',
        }

    def _parse_excel_file(self):
        """Parse the uploaded Excel file"""
        if not openpyxl:
            raise UserError(_('Please install openpyxl Python library to use this feature.'))

        if not self.file:
            raise UserError(_('Please upload an Excel file.'))

        try:
            file_content = base64.b64decode(self.file)
            wb = openpyxl.load_workbook(io.BytesIO(file_content))
            ws = wb.active

            # Get headers from first row
            headers = []
            for cell in ws[1]:
                headers.append(cell.value)

            # Validate required headers
            required_headers = ['lastname', 'name', 'id']
            for header in required_headers:
                if header not in headers:
                    raise UserError(_('Missing required column: %s') % header)

            # Parse data rows
            data = []
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not any(row):  # Skip empty rows
                    continue

                row_data = {}
                for idx, header in enumerate(headers):
                    if idx < len(row):
                        row_data[header] = row[idx]

                # Skip if required fields are empty
                if not row_data.get('lastname') or not row_data.get('name') or not row_data.get('id'):
                    continue

                data.append(row_data)

            return data

        except Exception as e:
            _logger.error(f"Error parsing Excel file: {e}")
            raise UserError(_('Error parsing Excel file: %s') % str(e))

    def _prepare_faculty_vals(self, row_data):
        """Prepare faculty values dictionary"""
        # Map gender to Odoo selection values
        gender_mapping = {
            'm': 'male',
            'male': 'male',
            'h': 'male',
            'f': 'female',
            'female': 'female',
            'o': 'other',
            'other': 'other',
        }
        
        lastname = row_data.get('lastname', '').strip()
        firstname = row_data.get('name', '').strip()
        vals = {
            'lastname': lastname,
            'firstname': firstname,
        }

        # Email
        if row_data.get('email'):
            vals['email'] = row_data.get('email', '').strip()

        # Gender
        if row_data.get('gender'):
            gender_value = gender_mapping.get(str(row_data['gender']).strip().lower(), False)
            if gender_value:
                vals['gender'] = gender_value

        # Second lastname
        if row_data.get('second_lastname'):
            lastname2 = row_data.get('second_lastname', '').strip()
            vals['lastname2'] = lastname2

        # ID/Document number
        if row_data.get('id'):
            vals['vat'] = row_data.get('id', '').strip()

        # Birthdate
        if row_data.get('birthdate'):
            try:
                if isinstance(row_data['birthdate'], datetime):
                    vals['birthdate_date'] = row_data['birthdate'].date()
                elif isinstance(row_data['birthdate'], str):
                    vals['birthdate_date'] = datetime.strptime(
                        row_data['birthdate'], '%Y-%m-%d'
                    ).date()
            except Exception as e:
                _logger.warning(f"Error parsing birthdate: {e}")

        # Speciality
        if row_data.get('speciality'):
            speciality_code = str(row_data['speciality']).strip()
            speciality = self.env['hezkuntza.speciality'].search([
                ('name', '=', speciality_code)
            ], limit=1)
            if speciality:
                vals['speciality_id'] = speciality.id
            else:
                _logger.warning(f"Speciality not found for code: {speciality_code}")
        vals["is_tutor"] = True
        return vals

    def action_import_faculties(self):
        """Import faculties from Excel file"""
        self.ensure_one()

        data = self._parse_excel_file()

        if not data:
            raise UserError(_('No valid data found in the Excel file.'))

        created_count = 0
        updated_count = 0
        error_count = 0
        errors = []

        for row_data in data:
            try:
                # Search for existing faculty by id_number or name
                faculty = None
                if row_data.get('id'):
                    faculty = self.env['res.partner'].search([
                        ('vat', '=', row_data['id'].strip())
                    ], limit=1)
                
                if not faculty:
                    # Search by name combination
                    domain = [
                        ('lastname', '=', row_data['lastname'].strip()),
                        ('name', '=', row_data['name'].strip())
                    ]
                    if row_data.get('second_lastname'):
                        domain.append(('lastname2', '=', row_data['second_lastname'].strip()))
                    
                    faculty = self.env['res.partner'].search(domain, limit=1)

                vals = self._prepare_faculty_vals(row_data)

                if faculty:
                    if self.override:
                        faculty.write(vals)
                        updated_count += 1
                    else:
                        _logger.info(f"Faculty already exists, skipping: {faculty.name}")
                else:
                    self.env['res.partner'].create(vals)
                    created_count += 1

            except Exception as e:
                error_count += 1
                error_msg = f"Row {row_data.get('name', 'Unknown')}: {str(e)}"
                errors.append(error_msg)
                _logger.error(f"Error importing faculty: {error_msg}")

        # Prepare result message
        message = _('Import completed!\n')
        message += _('Created: %s\n') % created_count
        message += _('Updated: %s\n') % updated_count
        if error_count > 0:
            message += _('Errors: %s\n') % error_count
            message += '\n'.join(errors[:10])  # Show first 10 errors
            if len(errors) > 10:
                message += _('\n... and %s more errors') % (len(errors) - 10)

        return {
            'type': 'ir.actions.client',
            'tag': 'display_notification',
            'params': {
                'title': _('Faculty Import'),
                'message': message,
                'type': 'success' if error_count == 0 else 'warning',
                'sticky': True,
                'next': {'type': 'ir.actions.act_window_close'},
            }
        }
